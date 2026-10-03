"""Generate Freightflex VM vs Cloud Run Word document."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"d:\Aniket\My WorkSpace\Professional\Server Code and DB\Flexishift\Freightflex\docs\Freightflex-VM-vs-Cloud-Run.docx"

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9)
    s.right_margin = Inches(0.9)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)
    return p


def add_para(text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run, size=11)


def shade_header_row(row):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "041627")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
                set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        hdr.cells[i].text = h
        for p in hdr.cells[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
    shade_header_row(hdr)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()
    return table


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Freightflex Deployment Decision")
set_run_font(r, size=20, bold=True, color=RGBColor(4, 22, 39))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("VM vs Cloud Run — Cost, Cold Start & Live Tracking")
set_run_font(r, size=12, color=RGBColor(70, 70, 70))

meta = doc.add_paragraph()
r = meta.add_run(
    "Project: Freightflex (FlexiShift)\n"
    "Context: ~100 users/month initially; Stripe payments + live GPS tracking (WebSockets)\n"
    "Goal: Choose a cost-effective host now, with a clear path to scale later"
)
set_run_font(r, size=10)

add_heading("1. Executive summary", 1)
add_para(
    "This document compares hosting Freightflex on a Virtual Machine (VM) versus Google Cloud Run, "
    "including cold starts, live tracking needs, and cost for low early traffic."
)

add_table(
    ["Approach", "Best for", "Approx cost (early)", "Cold start?"],
    [
        ["Small VM (4 GB RAM)", "Lowest cost now; simple all-in-one", "~$30–45 / month", "No (while VM is running)"],
        ["Cloud Run (scale-to-zero)", "Cheapest compute when idle", "~$30–70 / month (mostly DB)", "Yes, after idle"],
        ["Cloud Run (min-instances = 1)", "Prod live tracking without delay", "~$70–150 / month", "No"],
    ],
)

add_para("Recommendation for current stage (few clients, ~100 users/month):", bold=True)
add_bullets([
    "Prefer a 4 GB VM for lowest cost (not the previous 8 GB size).",
    "Move to Cloud Run later when traffic/clients grow.",
    "If using Cloud Run in real production with live tracking: set min-instances = 1.",
])

add_heading("2. What we actually need to run", 1)
add_para("Not everything on the old server must be redeployed. Only actively used services matter.")
add_table(
    ["Component", "Required?", "Notes"],
    [
        ["FastAPI backend (REST)", "Yes", "Core API"],
        ["Stripe webhooks", "Yes (payments)", "Same API URL / same service"],
        ["WebSockets (live tracking)", "Yes (feature exists)", "Needs a running instance during tracking"],
        ["MySQL", "Yes", "Application database"],
        ["File storage (docs/photos)", "Yes", "Local disk on VM, or object storage on Cloud Run"],
        ["Frontend (React)", "Yes", "Can use free static hosting"],
        ["Redis", "Optional", "App can fall back without it"],
        ["Celery", "Optional", "Skip unless required"],
        ["Load balancer", "Not now", "Add later if needed"],
    ],
)

add_heading("3. What is a cold start?", 1)
add_para(
    "Cold start means the delay when the platform must start your application from zero before it can handle a request."
)
add_para("Typical Cloud Run scale-to-zero flow:", bold=True)
add_bullets([
    "Nobody used the API for a while → 0 instances (asleep).",
    "User opens app / live map, or Stripe sends a webhook.",
    "Platform starts the container → cold start (about 1–5+ seconds).",
    "Request is handled; further requests are fast while the instance stays warm.",
    "Idle again → back to 0 → next user may hit another cold start.",
])

add_heading("3.1 Cold start by approach", 2)
add_table(
    ["Approach", "Cold start?", "Explanation"],
    [
        ["VM running 24×7", "No", "Process already running; connections are immediate."],
        [
            "VM stopped to save money, then started",
            "Yes (boot delay)",
            "OS + MySQL + API must start (can take minutes). Similar user wait, different mechanism.",
        ],
        ["Cloud Run scale-to-zero", "Yes", "After idle, first request wakes the service."],
        ["Cloud Run min-instances = 1", "No", "One instance always warm; higher idle cost."],
    ],
)
add_para(
    "Important: Choosing a VM does not create Cloud Run–style cold starts by itself. "
    "Cold starts appear if the service sleeps (Cloud Run scale-to-zero) or if the VM is powered off.",
    bold=True,
)

add_heading("4. Live tracking (WebSockets)", 1)
add_para(
    "Live tracking needs a running server while a trip is being tracked (open WebSocket). "
    "That does not automatically mean you must pay for always-on Cloud Run if you use a VM."
)
add_table(
    ["Need", "Meaning"],
    [
        ["During an active live job", "Server must be awake"],
        ["Instant map open anytime (no wait)", "Always-on: VM or Cloud Run min-instances = 1"],
        ["OK with occasional short delay", "Cloud Run scale-to-zero acceptable for early testing"],
    ],
)

add_heading("4.1 Same service or separate?", 2)
add_para(
    "Same service is correct for Freightflex today — one API process handles REST, Stripe webhooks, and WebSockets."
)
add_para(
    "One API process / one Cloud Run service / one VM app:\n"
    "  • REST API\n"
    "  • Stripe webhooks → same base URL\n"
    "  • WebSockets → same app (e.g. /ws/...)"
)
add_para("Split into a second service only later if tracking traffic grows a lot.")

add_heading("5. Approach A — Virtual Machine (VM)", 1)
add_heading("5.1 Architecture", 2)
add_para(
    "Users (Web / Mobile)\n"
    "        →\n"
    "   VM (e.g. 4 GB RAM)\n"
    "   • Nginx (optional)\n"
    "   • FastAPI (API + webhooks + WebSockets)\n"
    "   • MySQL\n"
    "   • Local files (/uploads) or object storage later\n"
    "Frontend: same VM or free static host"
)

add_heading("5.2 Pros", 2)
add_bullets([
    "Lowest cost for low traffic when rightsized (4 GB, not 8 GB).",
    "No Cloud Run–style cold start while the VM is on.",
    "Stripe + WebSockets on one box — simple.",
    "Familiar (matches previous hosting style).",
])

add_heading("5.3 Cons", 2)
add_bullets([
    "You manage OS updates, security, and backups.",
    "Pays 24×7 even at night with zero users.",
    "Harder to scale later than Cloud Run.",
    "If you stop the VM to save money → long boot delay (bad for webhooks/tracking).",
])

add_heading("5.4 Sizing guidance", 2)
add_table(
    ["RAM", "Approx cost", "Fit"],
    [
        ["8 GB (previous)", "~$60–70 / month", "Oversized for ~100 users/month"],
        ["4 GB (recommended now)", "~$30–45 / month", "Good for API + MySQL at low load"],
        ["2 GB", "~$15–30 / month", "Possible if tightly tuned; watch MySQL memory"],
    ],
)

add_heading("5.5 Cold start on VM?", 2)
add_bullets([
    "Running VM: no cold start.",
    "Stopped VM: yes — full machine boot (often worse UX than Cloud Run cold start).",
])

add_heading("6. Approach B — Cloud Run", 1)
add_heading("6.1 Architecture (no load balancer for now)", 2)
add_para(
    "Users (Web / Mobile)\n"
    "  • Frontend → free static hosting\n"
    "  • API → Cloud Run (single service)\n"
    "        – REST API\n"
    "        – Stripe webhooks\n"
    "        – WebSockets (live tracking)\n"
    "  • Cloud SQL (MySQL) → always-on cost\n"
    "  • Cloud Storage / Blob → files (not local disk)"
)

add_heading("6.2 Mode B1 — Scale-to-zero", 2)
add_para(
    "Scale-to-zero means: when nobody is using the app, Cloud Run shuts down all instances — "
    "you pay about $0 for compute while idle."
)
add_bullets([
    "Instances go to 0 when idle → compute ~$0 when unused.",
    "Has cold starts after idle.",
    "Fine for DEV / very early use.",
    "Live tracking / first API call after sleep may feel slow.",
    "Cloud SQL still costs even when the API is asleep.",
])

add_heading("6.3 Mode B2 — Min-instances = 1", 2)
add_para("Use this for production Cloud Run when live tracking must feel always ready.")
add_bullets([
    "Always one warm instance → no cold start.",
    "Required pattern if Cloud Run hosts real live-tracking UX.",
    "Costs more than scale-to-zero.",
    "Still one service for API + Stripe + WebSockets.",
])

add_heading("6.4 Pros / Cons", 2)
add_para("Pros:", bold=True)
add_bullets([
    "Less server ops than a VM.",
    "Auto scale when traffic grows.",
    "DEV can be cheap with scale-to-zero.",
    "Clean fit for containers + managed DB + object storage.",
])
add_para("Cons:", bold=True)
add_bullets([
    "Managed MySQL (Cloud SQL) is a major ongoing cost.",
    "Local /uploads disk does not work reliably — need object storage.",
    "Warm instance (min=1) reduces the “idle = free” benefit.",
    "Frontend free + Cloud Run free tier do not make the whole stack free.",
])

add_heading("6.5 Free tier reality", 2)
add_table(
    ["Piece", "Free?"],
    [
        ["Frontend (static)", "Often yes"],
        ["Cloud Run (low traffic, scale-to-zero)", "Free allowance can cover a lot"],
        ["Cloud Run min-instances = 1", "Not free (always allocated)"],
        ["Cloud SQL", "Not free (usually the main bill)"],
    ],
)

add_heading("7. Cost comparison (after free credits)", 1)
add_para(
    "Assumptions: ~100 users/month, Stripe + live tracking available, no load balancer. "
    "Prices vary by region — planning estimates only."
)
add_table(
    ["Setup", "Approx / month", "Cold start", "Live tracking UX"],
    [
        ["VM 8 GB (old size)", "~$60–70+", "No (if always on)", "Good"],
        ["VM 4 GB", "~$30–45", "No (if always on)", "Good"],
        ["Cloud Run scale-to-zero + small Cloud SQL", "~$30–70", "Yes after idle", "OK with occasional delay"],
        ["Cloud Run min-instances=1 + Cloud SQL", "~$70–150", "No", "Good"],
    ],
)
add_para(
    "During Google Cloud free trial credits (often ~$300 for ~90 days), much of this may be covered temporarily — "
    "confirm in your cloud console."
)

add_heading("8. Decision guide", 1)
add_para("Need lowest cost now, few users?", bold=True)
add_bullets([
    "→ VM 4 GB, always on",
    "→ No Cloud Run cold start while running",
    "→ Stripe + WebSockets on same VM",
])
add_para("Want managed scale later / less ops?", bold=True)
add_bullets([
    "→ Cloud Run",
    "→ Early / DEV: scale-to-zero (accept cold start)",
    "→ Real prod + live tracking: min-instances = 1 (no cold start)",
    "→ Cloud SQL + object storage",
    "→ Same Cloud Run service for API + Stripe + WebSockets",
])

add_heading("9. Final recommendation (current stage)", 1)
add_bullets([
    "Cost priority now: use a 4 GB VM (not 8 GB).",
    "Keep API + Stripe webhooks + WebSockets on that same host.",
    "Understand: VM running = no cold start; Cloud Run scale-to-zero = cold start; Cloud Run min-instances = 1 = no cold start.",
    "When clients grow: migrate to Cloud Run + Cloud SQL, with min-instances = 1 for live tracking.",
    "Skip load balancer, Redis, and Celery until needed.",
])

add_heading("10. One-line summary", 1)
add_para("VM (4 GB): cheapest now, always on, no cold start while running.", bold=True)
add_para("Cloud Run scale-to-zero: cheaper when idle, but cold starts can delay live tracking.")
add_para(
    "Cloud Run min-instances = 1: no cold start, better for live tracking, costs more — "
    "use when prod tracking matters on Cloud Run."
)

footer = doc.add_paragraph()
r = footer.add_run("\n— End of document —")
set_run_font(r, size=9, color=RGBColor(120, 120, 120))

doc.save(OUT)
print(OUT)
